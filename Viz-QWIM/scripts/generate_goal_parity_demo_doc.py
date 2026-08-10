"""Generate the QWIM Dashboard demo guide (.docx) with live screenshots.

Captures every main tab and Goal Parity subtab from a RUNNING dashboard via
headless Playwright (including one completed Historical Backtest run), then
assembles a client-facing walkthrough document. The output is committed as a
static asset served by the Overview & Guide subtab's download button.

Usage (dashboard must already be running):
    .venv/bin/shiny run src/dashboard/main_App.py &
    .venv/bin/python scripts/generate_goal_parity_demo_doc.py \
        [--url http://127.0.0.1:8000] [--out <path.docx>]
"""

from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from playwright.sync_api import sync_playwright


DEFAULT_OUT = (
    Path(__file__).resolve().parents[1]
    / "src" / "dashboard" / "shiny_tab_goal_parity" / "assets"
    / "QWIM_Dashboard_Demo_Guide.docx"
)

MAIN_TABS = ["Overview", "Setup", "Clients", "Portfolios", "Products", "Results"]
GP_SUBTABS = [
    "Overview & Guide",
    "Investor Profile",
    "4×4 Asset Map",
    "Strategic Optimization",
    "Tactical Rebalancing",
    "Historical Backtest",
]

IMAGE_WIDTH_IN = 6.3
#: Split full-page captures taller than this many px so no image overflows a page.
MAX_IMAGE_HEIGHT_PX = 1200


def _slug(name: str) -> str:
    return name.lower().replace(" & ", "_").replace(" ", "_").replace("×", "x")


def capture_screenshots(url: str, shots_dir: Path) -> None:
    """Screenshot every tab/subtab of the running dashboard."""
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": 1440, "height": 900})
        page.goto(url, wait_until="networkidle")
        page.wait_for_timeout(4000)

        for tab in MAIN_TABS:
            page.click(f"nav >> text={tab}")
            page.wait_for_timeout(2500)
            page.screenshot(path=shots_dir / f"tab_{_slug(tab)}.png")

        page.click("nav >> text=Goal Parity")
        page.wait_for_timeout(1500)
        for subtab in GP_SUBTABS:
            page.click(f"a:has-text('{subtab}')")
            page.wait_for_timeout(2500)
            page.screenshot(
                path=shots_dir / f"gp_{_slug(subtab)}.png",
                full_page=(subtab == "Overview & Guide"),
            )

        page.click("button:has-text('Run backtest')")
        page.wait_for_timeout(12000)
        page.screenshot(path=shots_dir / "gp_historical_backtest_results.png")
        browser.close()


def _add_picture(document: Document, path: Path, caption: str) -> None:
    """Insert an image (splitting page-overflowing captures) plus a caption."""
    image = Image.open(path)
    width, height = image.size
    if height > MAX_IMAGE_HEIGHT_PX:
        pieces = []
        for top in range(0, height, MAX_IMAGE_HEIGHT_PX):
            piece = image.crop((0, top, width, min(top + MAX_IMAGE_HEIGHT_PX, height)))
            piece_path = path.with_name(f"{path.stem}_part{top}.png")
            piece.save(piece_path)
            pieces.append(piece_path)
    else:
        pieces = [path]
    for piece_path in pieces:
        document.add_picture(str(piece_path), width=Inches(IMAGE_WIDTH_IN))
    paragraph = document.add_paragraph()
    run = paragraph.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6C, 0x75, 0x7D)


def build_document(shots_dir: Path, out_path: Path) -> None:
    document = Document()

    document.add_heading("QWIM Dashboard — Demo & User Guide", level=0)
    document.add_paragraph(
        "This guide walks through the QWIM dashboard the way we present it to "
        "business partners: what each tab shows, the order to demo it in, and "
        "the talking points for the Goal Parity investment model. Screenshots "
        "are taken from the live dashboard."
    )

    document.add_heading("Suggested demo flow", level=1)
    for line in [
        "Open on the Overview tab and state the value proposition in one line.",
        "Show Clients: the advisor-entered profile that personalizes everything downstream.",
        "Walk the Goal Parity tab left to right — it is a story in six steps.",
        "Close with the Historical Backtest: out-of-sample evidence, then the Results tab.",
    ]:
        document.add_paragraph(line, style="List Number")

    document.add_heading("The main tabs", level=1)
    tab_notes = {
        "Overview": (
            "The executive summary — the one-slide version of the QWIM "
            "approach. Start the demo here and keep it short."
        ),
        "Setup": (
            "Data sources, computation settings, and per-subtab console "
            "verbosity. In a partner demo you normally skip this tab; mention "
            "it exists so partners know configuration is centralized."
        ),
        "Clients": (
            "The advisor enters the client's personal information, assets, "
            "income, and goals here. Emphasize that models downstream — "
            "including Goal Parity — read this profile automatically: enter "
            "the client once, and every analysis is personalized."
        ),
        "Portfolios": (
            "Classical portfolio construction and comparison tooling "
            "(analysis, optimizers, benchmarks). Useful to frame Goal Parity "
            "against: this is the conventional toolkit."
        ),
        "Products": (
            "The investable product shelf backing the portfolios."
        ),
        "Results": (
            "Simulations and consolidated outputs, and where the client PDF "
            "report is produced. End the demo here: everything shown on "
            "screen can be handed to the client as a document."
        ),
    }
    for tab in MAIN_TABS:
        document.add_heading(tab, level=2)
        document.add_paragraph(tab_notes[tab])
        _add_picture(document, shots_dir / f"tab_{_slug(tab)}.png", f"The {tab} tab.")

    document.add_heading("The Goal Parity tab — step by step", level=1)
    document.add_paragraph(
        "Goal Parity builds the portfolio around four client goals — "
        "Liquidity, Income, Preservation, and Growth — instead of a single "
        "market benchmark. Demo the subtabs left to right; each step feeds "
        "the next."
    )
    gp_notes = {
        "Overview & Guide": (
            "The landing page: why goal-based investing holds up across "
            "market regimes, and the step-by-step usage guide. Partners can "
            "download this document from the button here."
        ),
        "Investor Profile": (
            "Step 1 — the client's risk comfort and horizon become the "
            "model's risk-aversion and loss barrier. Point out the 'Source:' "
            "line showing inputs flow in from the Clients tab."
        ),
        "4×4 Asset Map": (
            "Steps 2–4 — every asset is scored on how it actually behaves "
            "toward each goal (option-style triggers, not static labels). "
            "The scatter places each asset between the four goal corners."
        ),
        "Strategic Optimization": (
            "Step 5 — the optimizer balances the four goal powers (default: "
            "25% each) or applies a deliberate tilt. Show the weights and "
            "goal-power bars; note the warnings the page surfaces when "
            "something needs advisor judgement."
        ),
        "Tactical Rebalancing": (
            "Step 6 — a simulated market drift and the cost-aware trades "
            "that restore the goal balance. Emphasize the turnover budget: "
            "the model will not churn the portfolio."
        ),
        "Historical Backtest": (
            "The evidence: a walk-forward replay with no look-ahead — the "
            "model is re-trained on past data only at every step. The user "
            "chooses the date range, training window, test window, and step "
            "size, then runs it live."
        ),
    }
    for subtab in GP_SUBTABS:
        document.add_heading(subtab, level=2)
        document.add_paragraph(gp_notes[subtab])
        _add_picture(
            document,
            shots_dir / f"gp_{_slug(subtab)}.png",
            f"Goal Parity — {subtab}.",
        )

    document.add_heading("Backtest results (live run)", level=2)
    document.add_paragraph(
        "A completed run with the default settings (3-year training window, "
        "6-month step). The talking point: Goal Parity targets smaller "
        "drawdowns and lower volatility than the 60/40 benchmark — it gives "
        "up some upside in strong bull markets in exchange for a smoother "
        "ride through stress periods. It is not trying to beat the "
        "benchmark's raw return."
    )
    _add_picture(
        document,
        shots_dir / "gp_historical_backtest_results.png",
        "Historical Backtest after a completed run: out-of-sample equity "
        "curves and side-by-side performance metrics.",
    )

    document.add_heading("Practical notes for presenters", level=1)
    for note in [
        "Warning banners are features, not bugs: solver-failure and "
        "scarce-goal notes tell the advisor when judgement is needed — "
        "demo one deliberately if asked about model limits.",
        "Backtest guardrails reject impossible window configurations and "
        "warn on training windows under ~2.5 years.",
        "All figures are illustrative and do not constitute investment "
        "advice.",
    ]:
        document.add_paragraph(note, style="List Bullet")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    print(f"Wrote {out_path} ({out_path.stat().st_size / 1e6:.1f} MB)")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:8000")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory() as tmp:
        shots_dir = Path(tmp)
        capture_screenshots(args.url, shots_dir)
        build_document(shots_dir, args.out)


if __name__ == "__main__":
    main()
